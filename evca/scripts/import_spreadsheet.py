#!/usr/bin/env python3
"""
import_spreadsheet.py  —  EVCA spreadsheet → SQLite DB

Usage:
    python3 import_spreadsheet.py --load-id N --db PATH --spreadsheet PATH

Phase 1: structural validation against metadata.sql
Phase 2: tab-by-tab data extraction and INSERT (dependency order)

Exit 0 = success (or success with warnings), 1 = failure
"""

import argparse
import os
import sqlite3
import sys
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

SCRIPT_DIR   = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR  = os.path.dirname(SCRIPT_DIR)
METADATA_SQL = os.path.join(PROJECT_DIR, 'output', 'metadata.sql')

EXCEL_ERRORS = frozenset({'#N/A', '#REF!', '#VALUE!', '#DIV/0!', '#NULL!', '#NAME?', '#NUM!'})

# Per-hazard blocks: column to check for data presence (used in validation)
HAZARD_DATA_COL = {22: 'G', 23: 'G', 24: 'G', 29: 'M', 30: 'M', 31: 'M'}


# ── report collector ──────────────────────────────────────────────────────────

class Report:
    def __init__(self):
        self.lines = []
        self.errors = 0
        self.warnings = 0

    def info(self, msg):    self.lines.append(f'INFO  {msg}')
    def warn(self, msg):    self.warnings += 1; self.lines.append(f'WARN  {msg}')
    def error(self, msg):   self.errors += 1;   self.lines.append(f'ERROR {msg}')
    def section(self, msg): self.lines.extend(['', f'── {msg}', ''])
    def text(self):         return '\n'.join(self.lines)


# ── cell helpers ──────────────────────────────────────────────────────────────

def cv(ws, col_letter, row_num):
    """Return stripped string cell value, or None if blank/error."""
    col = column_index_from_string(col_letter)
    v = ws.cell(row=row_num, column=col).value
    if v is None:
        return None
    s = str(v).strip()
    return None if (not s or s in EXCEL_ERRORS) else s


def cv_raw(ws, col_letter, row_num):
    """Return raw value (preserves int/float), or None if blank/error."""
    col = column_index_from_string(col_letter)
    v = ws.cell(row=row_num, column=col).value
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    return None if (not s or s in EXCEL_ERRORS) else v


def strip_pfx(val):
    """Strip 'Label: ' prefix from labelled cell values."""
    if not val:
        return None
    s = str(val).strip()
    if ': ' in s:
        return s.split(': ', 1)[1].strip() or None
    return s or None


def parse_dim_number(label):
    """'3. Air dan Sanitasi' → 3.  Returns None if unparseable."""
    if not label:
        return None
    try:
        return int(str(label).strip().split('.')[0].strip())
    except (ValueError, IndexError):
        return None


def to_int(v):
    if v is None:
        return None
    try:
        return int(v)
    except (ValueError, TypeError):
        return None


def to_real(v):
    if v is None:
        return None
    try:
        return float(v)
    except (ValueError, TypeError):
        return None


# ── metadata loader ───────────────────────────────────────────────────────────

def load_metadata(path):
    conn = sqlite3.connect(':memory:')
    conn.row_factory = sqlite3.Row
    with open(path, encoding='utf-8') as f:
        conn.executescript(f.read())
    return conn


# ── Phase 1: structural validation ───────────────────────────────────────────

def _count_nonempty(ws, col_letter, row_start, row_end):
    col = column_index_from_string(col_letter)
    return sum(
        1 for r in range(row_start, row_end + 1)
        if (lambda v: v is not None and str(v).strip() and str(v).strip() not in EXCEL_ERRORS)(
            ws.cell(row=r, column=col).value
        )
    )


def _validate_block(block, ws, meta_conn):
    bid       = block['id']
    col_start = block['col_start']
    row_start = block['row_start']
    row_end   = block['row_end']
    btype     = block['block_type']
    bname     = block['block_name_english'] or block['block_name_original'] or f'block {bid}'
    is_h3     = 'Hazard 3' in (bname or '')
    details   = []

    if btype == 'reference':
        expected = meta_conn.execute(
            'SELECT COUNT(*) FROM evca_reference_data WHERE block_id=?', (bid,)
        ).fetchone()[0]
        actual = _count_nonempty(ws, col_start, row_start, row_end)
        details.append(f'ref rows: {actual}/{expected}')
        if actual == 0:
            return 'EMPTY', details
        return ('WARN!' if actual < expected else 'PASS'), details

    if bid in HAZARD_DATA_COL:
        dc = HAZARD_DATA_COL[bid]
        details.append(f'rows {row_start}–{row_end}')
        first = cv(ws, dc, row_start)
        last  = cv(ws, dc, row_end)
        if first is not None:
            details.append(f'{dc}{row_start}="{str(first)[:25]}"')
            return 'PASS', details
        if last is not None:
            details.append(f'{dc}{row_start}=empty {dc}{row_end}="{str(last)[:20]}" (social dims only)')
            return 'PASS', details
        if is_h3:
            details.append('all empty (expected — 2-hazard assessment)')
            return 'EMPTY', details
        details.append('all empty')
        return 'WARN!', details

    anchor = cv(ws, col_start, row_start)
    details.append(f'anchor {col_start}{row_start}={"OK" if anchor else "empty"}')
    fields = meta_conn.execute(
        'SELECT field_name_english, column_letter, row_number FROM evca_block_fields '
        'WHERE block_id=? AND column_letter IS NOT NULL AND row_number IS NOT NULL',
        (bid,)
    ).fetchall()
    hits = sum(1 for f in fields if cv(ws, f['column_letter'], f['row_number']) is not None)
    if fields:
        details.append(f'fields {hits}/{len(fields)}')

    if anchor is not None or hits > 0:
        return 'PASS', details
    if is_h3:
        return 'EMPTY', details
    return 'WARN!', details


def run_validation(wb, meta_conn, report):
    """Returns True if safe to proceed with import."""
    report.section('Phase 1: Structural Validation')
    sheets = wb.worksheets
    blocks = meta_conn.execute('SELECT * FROM evca_sheet_blocks ORDER BY id').fetchall()
    counts = {'PASS': 0, 'WARN!': 0, 'EMPTY': 0}
    critical_warn = False
    cur_sheet = None

    for block in blocks:
        si = block['sheet_index']
        if si != cur_sheet:
            cur_sheet = si
            report.info(f'Sheet {si}: {block["sheet_name_english"]}')
        ws = sheets[si]
        status, details = _validate_block(block, ws, meta_conn)
        counts[status] += 1
        bname = block['block_name_english'] or block['block_name_original'] or f'block {block["id"]}'
        rng   = f'{block["cell_range_start"]}:{block["cell_range_end"]}'
        msg   = f'[{status}] #{block["id"]:02d} {bname} ({rng}) — {" | ".join(details)}'
        if status == 'WARN!':
            report.warn(msg)
            if block['id'] in (1, 14):   # General Info and Hazard Data are critical
                critical_warn = True
                report.error(f'Critical block #{block["id"]} failed — import cannot proceed')
        else:
            report.info(msg)

    report.info(f'Validation: PASS={counts["PASS"]} WARN={counts["WARN!"]} EMPTY={counts["EMPTY"]}')
    if critical_warn:
        report.error('Import aborted due to critical validation failures.')
        return False
    if counts['WARN!'] > 0:
        report.warn(f'{counts["WARN!"]} non-critical block warning(s) — proceeding')
    return True


# ── Phase 2: tab-by-tab import ────────────────────────────────────────────────

def import_tab1_assessment(ws, db_conn, load_id, report):
    """Block 1 + Block 2 from Tab 1.  Returns assessment_id or None on failure."""
    report.section('Tab 1: Assessment / Summary (General Info + Narrative)')

    village_name = cv(ws, 'E', 6)
    if not village_name:
        report.error('Tab 1: village_name (E6) is empty — cannot create assessment record')
        return None

    try:
        cur = db_conn.execute(
            '''INSERT INTO evca_assessments
               (load_id, village_name, district, province, country, national_society,
                branch, start_date, responsible_person, end_date, responsible_position,
                gps_coordinates, email, assessment_selection_narrative)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
            (load_id, village_name,
             strip_pfx(cv(ws, 'G', 6)),                    # district
             strip_pfx(cv(ws, 'E', 8)),                    # province
             strip_pfx(cv(ws, 'G', 8)) or 'Indonesia',     # country
             strip_pfx(cv(ws, 'E', 10)) or 'PMI',          # national_society
             strip_pfx(cv(ws, 'G', 10)),                   # branch
             strip_pfx(cv(ws, 'E', 12)),                   # start_date
             strip_pfx(cv(ws, 'G', 12)),                   # responsible_person
             strip_pfx(cv(ws, 'E', 14)),                   # end_date
             strip_pfx(cv(ws, 'G', 14)),                   # responsible_position
             strip_pfx(cv(ws, 'E', 16)),                   # gps_coordinates
             strip_pfx(cv(ws, 'G', 16)),                   # email
             cv(ws, 'C', 24),                              # assessment_selection_narrative
             )
        )
        db_conn.commit()
        aid = cur.lastrowid
        report.info(f'Tab 1: assessment inserted id={aid} village="{village_name}"')
        return aid
    except Exception as e:
        db_conn.rollback()
        report.error(f'Tab 1: assessment INSERT failed: {e}')
        return None


def import_tab2(ws, db_conn, load_id, assessment_id, report):
    """Blocks 5–8: community_context + population."""
    report.section('Tab 2: Background / Community Context')

    try:
        db_conn.execute(
            '''INSERT INTO evca_community_context
               (load_id, assessment_id, community_description,
                evca_participants_male, evca_participants_female,
                community_type, community_type_comments,
                geophysical_env_1, geophysical_env_2, geophysical_comments,
                livelihood_primary, livelihood_secondary, livelihood_other,
                assessment_process)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
            (load_id, assessment_id,
             cv(ws, 'C', 13),                       # community_description
             to_int(cv_raw(ws, 'U', 26)),            # evca_participants_male
             to_int(cv_raw(ws, 'V', 26)),            # evca_participants_female
             cv(ws, 'E', 29),                        # community_type
             cv(ws, 'E', 33),                        # community_type_comments
             cv(ws, 'M', 29),                        # geophysical_env_1
             cv(ws, 'M', 31),                        # geophysical_env_2
             cv(ws, 'M', 33),                        # geophysical_comments
             cv(ws, 'U', 29),                        # livelihood_primary
             cv(ws, 'U', 31),                        # livelihood_secondary
             cv(ws, 'U', 33),                        # livelihood_other
             cv(ws, 'C', 42),                        # assessment_process
             )
        )
        db_conn.commit()
        report.info('Tab 2: community_context inserted')
    except Exception as e:
        db_conn.rollback()
        report.error(f'Tab 2: community_context INSERT failed: {e}')

    # Population table: row 20 = counts, row 24 = disability
    age_groups = [('0_5', 'E', 'F'), ('6_17', 'I', 'J'), ('18_65', 'M', 'N'), ('66_plus', 'Q', 'R')]
    pop_count = 0
    for age_group, mc, fc in age_groups:
        m = to_int(cv_raw(ws, mc, 20))
        f = to_int(cv_raw(ws, fc, 20))
        if m is None and f is None:
            continue
        try:
            db_conn.execute(
                '''INSERT INTO evca_population
                   (load_id, assessment_id, age_group, male_count, female_count,
                    disability_male, disability_female)
                   VALUES (?,?,?,?,?,?,?)''',
                (load_id, assessment_id, age_group, m, f,
                 to_int(cv_raw(ws, mc, 24)) or 0,
                 to_int(cv_raw(ws, fc, 24)) or 0)
            )
            pop_count += 1
        except Exception as e:
            report.error(f'Tab 2: population {age_group}: {e}')
    db_conn.commit()
    report.info(f'Tab 2: {pop_count} population rows inserted')


def import_tab3_hazards(ws, db_conn, load_id, assessment_id, report):
    """Blocks 13–15 + canonical names from col AA.  Returns {hazard_number: db_id}."""
    report.section('Tab 3: Priority Hazards')

    # Update assessment with Tab 3 fields
    try:
        db_conn.execute(
            'UPDATE evca_assessments SET hazard_selection_rationale=?, priority_hazard_count=? WHERE id=?',
            (cv(ws, 'C', 14), to_int(cv_raw(ws, 'V', 48)), assessment_id)
        )
        db_conn.commit()
        report.info(f'Tab 3: assessment updated (hazard_selection_rationale, priority_count={to_int(cv_raw(ws, "V", 48))})')
    except Exception as e:
        db_conn.rollback()
        report.error(f'Tab 3: assessment UPDATE failed: {e}')

    # H1 (cols C-K, rows 16-30), H2 (cols M-W, rows 16-30), H3 (cols C-K, rows 32-46)
    # Canonical names from col AA (formula-derived); fall back to cell value
    hazard_defs = [
        (1, cv(ws, 'AA', 2) or cv(ws, 'E', 18),
         {'cause_origin': ('E',20), 'warning_signs': ('E',22), 'action_time': ('E',24),
          'frequency': ('E',26), 'occurrence_period': ('E',28), 'duration': ('E',30)}),
        (2, cv(ws, 'AA', 3) or cv(ws, 'O', 18),
         {'cause_origin': ('O',20), 'warning_signs': ('O',22), 'action_time': ('O',24),
          'frequency': ('O',26), 'occurrence_period': ('O',28), 'duration': ('O',30)}),
        (3, cv(ws, 'E', 34),
         {'cause_origin': ('E',36), 'warning_signs': ('E',38), 'action_time': ('E',40),
          'frequency': ('E',42), 'occurrence_period': ('E',44), 'duration': ('E',46)}),
    ]

    hazard_ids = {}
    for hnum, hname, field_map in hazard_defs:
        if not hname:
            report.info(f'Tab 3: Hazard {hnum} has no name — skipping')
            continue
        try:
            cur = db_conn.execute(
                '''INSERT INTO evca_hazards
                   (load_id, assessment_id, hazard_number, hazard_name,
                    cause_origin, warning_signs, action_time,
                    frequency, occurrence_period, duration)
                   VALUES (?,?,?,?,?,?,?,?,?,?)''',
                (load_id, assessment_id, hnum, hname,
                 cv(ws, *field_map['cause_origin']),
                 cv(ws, *field_map['warning_signs']),
                 cv(ws, *field_map['action_time']),
                 cv(ws, *field_map['frequency']),
                 cv(ws, *field_map['occurrence_period']),
                 cv(ws, *field_map['duration']))
            )
            db_conn.commit()
            hazard_ids[hnum] = cur.lastrowid
            report.info(f'Tab 3: Hazard {hnum} "{hname}" inserted id={cur.lastrowid}')
        except Exception as e:
            db_conn.rollback()
            report.error(f'Tab 3: Hazard {hnum} INSERT failed: {e}')

    return hazard_ids


def import_tab1_risk_summary(ws0, db_conn, load_id, assessment_id, hazard_ids, report):
    """Block 3 from Tab 1: risk summary per hazard × dimension (derived from Sheet 7)."""
    report.section('Tab 1: Risk Pattern Summary (Block 3)')
    # H1: rows 41-51, H2: rows 55-65, H3: rows 69-79; risk_label in col E
    hazard_row_ranges = {1: (41, 51), 2: (55, 65), 3: (69, 79)}
    count = 0
    for hnum, (row_start, row_end) in hazard_row_ranges.items():
        hazard_id = hazard_ids.get(hnum)
        if not hazard_id:
            report.info(f'Tab 1 risk summary: Hazard {hnum} not in DB — skipping')
            continue
        for dim_offset, row in enumerate(range(row_start, row_end + 1)):
            risk_label = cv(ws0, 'E', row)
            if risk_label is None:
                continue
            try:
                db_conn.execute(
                    '''INSERT OR IGNORE INTO evca_risk_summary
                       (load_id, assessment_id, hazard_id, dimension_number, risk_label)
                       VALUES (?,?,?,?,?)''',
                    (load_id, assessment_id, hazard_id, dim_offset + 1, risk_label)
                )
                count += 1
            except Exception as e:
                report.error(f'Tab 1 risk summary H{hnum} dim {dim_offset+1}: {e}')
    db_conn.commit()
    report.info(f'Tab 1: {count} risk_summary rows inserted')


def import_tab4(ws, db_conn, load_id, assessment_id, hazard_ids, report):
    """Blocks 17–19: vulnerability overview, vulnerable groups, vulnerability ratings."""
    report.section('Tab 4: Vulnerability')

    try:
        db_conn.execute(
            'UPDATE evca_assessments SET vulnerability_overview=? WHERE id=?',
            (cv(ws, 'C', 12), assessment_id)
        )
        db_conn.commit()
    except Exception as e:
        report.error(f'Tab 4: vulnerability_overview UPDATE failed: {e}')

    # Vulnerable groups: data rows 16, 18, 20, 22
    group_count = 0
    for gnum, row in enumerate([16, 18, 20, 22], 1):
        name = cv(ws, 'E', row)
        if not name:
            continue
        try:
            db_conn.execute(
                '''INSERT INTO evca_vulnerable_groups
                   (load_id, assessment_id, group_number, group_name, vulnerability_reasons)
                   VALUES (?,?,?,?,?)''',
                (load_id, assessment_id, gnum, name, cv(ws, 'O', row))
            )
            group_count += 1
        except Exception as e:
            report.error(f'Tab 4: group {gnum}: {e}')
    db_conn.commit()
    report.info(f'Tab 4: {group_count} vulnerable_groups inserted')

    # Vulnerability ratings: H1 rows 29-36, H2 rows 43-50, H3 rows 57-64
    hazard_row_ranges = {1: (29, 36), 2: (43, 50), 3: (57, 64)}
    rating_count = 0
    for hnum, (row_start, row_end) in hazard_row_ranges.items():
        hazard_id = hazard_ids.get(hnum)
        if not hazard_id:
            continue
        for row in range(row_start, row_end + 1):
            dim_number = parse_dim_number(cv(ws, 'C', row))
            if dim_number is None:
                continue
            try:
                db_conn.execute(
                    '''INSERT OR IGNORE INTO evca_vulnerability_ratings
                       (load_id, assessment_id, hazard_id, dimension_number,
                        impact_description, vulnerability_aspects, rating_label, rating_value)
                       VALUES (?,?,?,?,?,?,?,?)''',
                    (load_id, assessment_id, hazard_id, dim_number,
                     cv(ws, 'G', row), cv(ws, 'M', row),
                     cv(ws, 'S', row), to_real(cv_raw(ws, 'V', row)))
                )
                rating_count += 1
            except Exception as e:
                report.error(f'Tab 4: vuln rating H{hnum} dim {dim_number}: {e}')
    db_conn.commit()
    report.info(f'Tab 4: {rating_count} vulnerability_ratings inserted')


def import_tab5(ws, db_conn, load_id, assessment_id, hazard_ids, report):
    """Blocks 21–24: capacity overview + capacity ratings."""
    report.section('Tab 5: Capacity')

    try:
        db_conn.execute(
            'UPDATE evca_assessments SET capacity_overview=? WHERE id=?',
            (cv(ws, 'C', 14), assessment_id)
        )
        db_conn.commit()
    except Exception as e:
        report.error(f'Tab 5: capacity_overview UPDATE failed: {e}')

    # H1: rows 19-28, H2: rows 33-42, H3: rows 47-56
    hazard_row_ranges = {1: (19, 28), 2: (33, 42), 3: (47, 56)}
    count = 0
    for hnum, (row_start, row_end) in hazard_row_ranges.items():
        hazard_id = hazard_ids.get(hnum)
        if not hazard_id:
            continue
        for row in range(row_start, row_end + 1):
            dim_number = parse_dim_number(cv(ws, 'C', row))
            if dim_number is None:
                continue    # #REF! rows (dims 9-10) and spacers skipped here
            try:
                db_conn.execute(
                    '''INSERT OR IGNORE INTO evca_capacity_ratings
                       (load_id, assessment_id, hazard_id, dimension_number,
                        capacity_description, rating_label, rating_value)
                       VALUES (?,?,?,?,?,?,?)''',
                    (load_id, assessment_id, hazard_id, dim_number,
                     cv(ws, 'G', row), cv(ws, 'M', row),
                     to_real(cv_raw(ws, 'P', row)))
                )
                count += 1
            except Exception as e:
                report.error(f'Tab 5: cap rating H{hnum} dim {dim_number}: {e}')
    db_conn.commit()
    report.info(f'Tab 5: {count} capacity_ratings inserted')


def import_tab6(ws, db_conn, load_id, assessment_id, report):
    """Blocks 26–27: social dimensions overview + ratings."""
    report.section('Tab 6: Social Cohesion, Inclusion & Connectedness')

    try:
        db_conn.execute(
            'UPDATE evca_assessments SET social_dimensions_overview=? WHERE id=?',
            (cv(ws, 'C', 12), assessment_id)
        )
        db_conn.commit()
    except Exception as e:
        report.error(f'Tab 6: social_dimensions_overview UPDATE failed: {e}')

    dim_keys = {16: 'social_cohesion', 17: 'inclusion', 18: 'connectedness'}
    count = 0
    for row, dim_key in dim_keys.items():
        try:
            db_conn.execute(
                '''INSERT OR IGNORE INTO evca_social_dimensions
                   (load_id, assessment_id, dimension, description, rating_label, rating_value)
                   VALUES (?,?,?,?,?,?)''',
                (load_id, assessment_id, dim_key,
                 cv(ws, 'G', row), cv(ws, 'S', row),
                 to_real(cv_raw(ws, 'T', row)))
            )
            count += 1
        except Exception as e:
            report.error(f'Tab 6: social dim {dim_key}: {e}')
    db_conn.commit()
    report.info(f'Tab 6: {count} social_dimensions inserted')


def import_tab7(ws, db_conn, load_id, assessment_id, hazard_ids, report):
    """Blocks 29–31: risk ratings (all computed, col M = risk_label)."""
    report.section('Tab 7: Risk Ratings')

    # H1: rows 15-25, H2: rows 30-40, H3: rows 45-55 (11 rows = 11 dimensions each)
    hazard_row_ranges = {1: (15, 25), 2: (30, 40), 3: (45, 55)}
    count = 0
    for hnum, (row_start, row_end) in hazard_row_ranges.items():
        hazard_id = hazard_ids.get(hnum)
        if not hazard_id:
            continue
        for dim_offset, row in enumerate(range(row_start, row_end + 1)):
            risk_label = cv(ws, 'M', row)
            if risk_label is None:
                continue
            try:
                db_conn.execute(
                    '''INSERT OR IGNORE INTO evca_risk_ratings
                       (load_id, assessment_id, hazard_id, dimension_number, risk_label)
                       VALUES (?,?,?,?,?)''',
                    (load_id, assessment_id, hazard_id, dim_offset + 1, risk_label)
                )
                count += 1
            except Exception as e:
                report.error(f'Tab 7: risk rating H{hnum} dim {dim_offset+1}: {e}')
    db_conn.commit()
    report.info(f'Tab 7: {count} risk_ratings inserted')


def import_tab8(ws, db_conn, load_id, assessment_id, hazard_ids, report):
    """Blocks 32–33: risk analysis table + overall narrative."""
    report.section('Tab 8: Risk Analysis')

    # Col C has hazard name on first row of each group (merged cell; None for subsequent rows).
    # Use positional ordering: first hazard name found = H1, second = H2, etc.
    ordered_hids = [hazard_ids[k] for k in sorted(hazard_ids)]
    cur_hazard_id = None
    hazard_idx = -1
    count = 0

    for row in range(14, 60):
        hazard_name_cell = cv(ws, 'C', row)
        if hazard_name_cell:
            hazard_idx += 1
            cur_hazard_id = ordered_hids[hazard_idx] if hazard_idx < len(ordered_hids) else None
            if cur_hazard_id is None:
                report.warn(f'Tab 8: row {row} extra hazard "{hazard_name_cell}" — no matching DB record')

        dim_cell = cv(ws, 'G', row)
        dim_number = parse_dim_number(dim_cell)
        if dim_number is None or cur_hazard_id is None:
            continue

        try:
            db_conn.execute(
                '''INSERT OR IGNORE INTO evca_risk_analysis
                   (load_id, assessment_id, hazard_id, consolidated_dimension,
                    vulnerability_aspects, capacity_aspects, key_risk_summary)
                   VALUES (?,?,?,?,?,?,?)''',
                (load_id, assessment_id, cur_hazard_id, dim_number,
                 cv(ws, 'J', row), cv(ws, 'L', row), cv(ws, 'N', row))
            )
            count += 1
        except Exception as e:
            report.error(f'Tab 8: risk_analysis row {row}: {e}')

    db_conn.commit()
    report.info(f'Tab 8: {count} risk_analysis rows inserted')

    narrative = cv(ws, 'C', 30)
    if narrative:
        try:
            db_conn.execute(
                '''INSERT OR IGNORE INTO evca_overall_analysis
                   (load_id, assessment_id, overall_narrative)
                   VALUES (?,?,?)''',
                (load_id, assessment_id, narrative)
            )
            db_conn.commit()
            report.info('Tab 8: overall_analysis inserted')
        except Exception as e:
            db_conn.rollback()
            report.error(f'Tab 8: overall_analysis INSERT failed: {e}')


def import_tab9(ws, db_conn, load_id, assessment_id, report):
    """Blocks 35–36: action plan items + validation signatures."""
    report.section('Tab 9: Action Plan')

    # Data rows at 13,15,17,19,21,23,25 (step=2); rows 27+ are empty placeholders
    item_count = 0
    for item_order, row in enumerate(range(13, 32, 2), 1):
        risk_desc = cv(ws, 'C', row)
        if not risk_desc:
            break
        try:
            db_conn.execute(
                '''INSERT INTO evca_action_items
                   (load_id, assessment_id, item_order, priority_risk_description,
                    desired_outcome, priority_activities, required_resources,
                    technical_support_needs, schedule, responsible_party)
                   VALUES (?,?,?,?,?,?,?,?,?,?)''',
                (load_id, assessment_id, item_order, risk_desc,
                 cv(ws, 'E', row), cv(ws, 'G', row), cv(ws, 'I', row),
                 cv(ws, 'K', row), cv(ws, 'M', row), cv(ws, 'O', row))
            )
            item_count += 1
        except Exception as e:
            report.error(f'Tab 9: action_item {item_order}: {e}')
    db_conn.commit()
    report.info(f'Tab 9: {item_count} action_items inserted')

    try:
        db_conn.execute(
            '''INSERT OR IGNORE INTO evca_action_validation
               (load_id, assessment_id, community_rep_name, bpbd_rep_name,
                village_rep_name, pmi_rep_name)
               VALUES (?,?,?,?,?,?)''',
            (load_id, assessment_id,
             cv(ws, 'C', 36), cv(ws, 'E', 36), cv(ws, 'G', 36), cv(ws, 'I', 36))
        )
        db_conn.commit()
        report.info('Tab 9: action_validation inserted')
    except Exception as e:
        db_conn.rollback()
        report.error(f'Tab 9: action_validation INSERT failed: {e}')


def import_tab10(ws, db_conn, load_id, assessment_id, report):
    """Block 37: priority scoring matrix — column-major layout."""
    report.section('Tab 10: Priority Scoring Matrix')

    # Activity names in row 2, cols C–T.  Scores in rows 4–11 of same column.
    score_row_map = {
        4: 'score_funding',            5: 'score_timeframe',
        6: 'score_local_resources',    7: 'score_community_participation',
        8: 'score_govt_support',       9: 'score_sustainability',
        10: 'score_pmi_mandate',       11: 'score_effectiveness',
    }
    count = 0
    for col_letter in [chr(c) for c in range(ord('C'), ord('T') + 1)]:
        activity_name = cv(ws, col_letter, 2)
        if not activity_name:
            continue
        scores = {col: to_int(cv_raw(ws, col_letter, row)) for row, col in score_row_map.items()}
        try:
            db_conn.execute(
                '''INSERT INTO evca_priority_scores_flat
                   (load_id, assessment_id, activity_name,
                    score_funding, score_timeframe, score_local_resources,
                    score_community_participation, score_govt_support,
                    score_sustainability, score_pmi_mandate, score_effectiveness)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?)''',
                (load_id, assessment_id, activity_name,
                 scores['score_funding'], scores['score_timeframe'], scores['score_local_resources'],
                 scores['score_community_participation'], scores['score_govt_support'],
                 scores['score_sustainability'], scores['score_pmi_mandate'], scores['score_effectiveness'])
            )
            count += 1
        except Exception as e:
            report.error(f'Tab 10: "{activity_name}": {e}')
    db_conn.commit()
    report.info(f'Tab 10: {count} priority_scores inserted')


# ── main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description='Import EVCA spreadsheet into SQLite DB')
    parser.add_argument('--load-id',     required=True, type=int)
    parser.add_argument('--db',          required=True)
    parser.add_argument('--spreadsheet', required=True)
    args = parser.parse_args()

    report = Report()
    report.info(f'Started: {datetime.now().isoformat()}')
    report.info(f'Spreadsheet: {os.path.basename(args.spreadsheet)}')
    report.info(f'Load ID: {args.load_id}')

    for label, path in [('Metadata', METADATA_SQL), ('Spreadsheet', args.spreadsheet), ('Database', args.db)]:
        if not os.path.exists(path):
            report.error(f'{label} not found: {path}')
            print(report.text())
            sys.exit(1)

    db_conn = sqlite3.connect(args.db, timeout=30)
    db_conn.row_factory = sqlite3.Row
    db_conn.execute('PRAGMA foreign_keys = ON')
    db_conn.execute('PRAGMA journal_mode = WAL')
    db_conn.execute(
        "UPDATE evca_spreadsheet_loads SET status='processing' WHERE id=?", (args.load_id,)
    )
    db_conn.commit()

    meta_conn = load_metadata(METADATA_SQL)
    wb = load_workbook(args.spreadsheet, data_only=True)

    try:
        if not run_validation(wb, meta_conn, report):
            db_conn.execute(
                "UPDATE evca_spreadsheet_loads SET status='failure', report=? WHERE id=?",
                (report.text(), args.load_id)
            )
            db_conn.commit()
            print(report.text())
            sys.exit(1)

        sheets = wb.worksheets

        assessment_id = import_tab1_assessment(sheets[0], db_conn, args.load_id, report)
        if assessment_id is None:
            report.error('Cannot continue without assessment record.')
            db_conn.execute(
                "UPDATE evca_spreadsheet_loads SET status='failure', report=? WHERE id=?",
                (report.text(), args.load_id)
            )
            db_conn.commit()
            print(report.text())
            sys.exit(1)

        import_tab2(sheets[1], db_conn, args.load_id, assessment_id, report)

        hazard_ids = import_tab3_hazards(sheets[2], db_conn, args.load_id, assessment_id, report)
        if not hazard_ids:
            report.error('No hazards inserted — all hazard-dependent tabs will be skipped')

        import_tab1_risk_summary(sheets[0], db_conn, args.load_id, assessment_id, hazard_ids, report)
        import_tab4(sheets[3], db_conn, args.load_id, assessment_id, hazard_ids, report)
        import_tab5(sheets[4], db_conn, args.load_id, assessment_id, hazard_ids, report)
        import_tab6(sheets[5], db_conn, args.load_id, assessment_id, report)
        import_tab7(sheets[6], db_conn, args.load_id, assessment_id, hazard_ids, report)
        import_tab8(sheets[7], db_conn, args.load_id, assessment_id, hazard_ids, report)
        import_tab9(sheets[8], db_conn, args.load_id, assessment_id, report)
        import_tab10(sheets[9], db_conn, args.load_id, assessment_id, report)

        status = 'failure' if report.errors > 0 else 'success'
        report.section('Import Complete')
        report.info(f'Status: {status}  |  Errors: {report.errors}  Warnings: {report.warnings}')
        report.info(f'Finished: {datetime.now().isoformat()}')

        db_conn.execute(
            "UPDATE evca_spreadsheet_loads SET status=?, report=? WHERE id=?",
            (status, report.text(), args.load_id)
        )
        db_conn.commit()

    except Exception as e:
        import traceback
        report.error(f'Unexpected error: {e}\n{traceback.format_exc()}')
        db_conn.execute(
            "UPDATE evca_spreadsheet_loads SET status='failure', report=? WHERE id=?",
            (report.text(), args.load_id)
        )
        db_conn.commit()
        print(report.text(), file=sys.stderr)
        sys.exit(1)
    finally:
        wb.close()
        meta_conn.close()
        db_conn.close()

    print(report.text())
    sys.exit(0 if report.errors == 0 else 1)


if __name__ == '__main__':
    main()
